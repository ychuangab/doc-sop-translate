"""產生範例 SOP（中文 DOCX）。內容全為虛構。執行：python3 samples-src/make_sample.py"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "public" / "samples" / "SOP_空壓機日常點檢.docx"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Microsoft JhengHei"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
st.font.size = Pt(11)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def h(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12.5)


title("空壓機日常點檢標準作業程序")
p = doc.add_paragraph("文件編號：SOP-MT-012　版次：A　（範例文件，內容為虛構）")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

h("1. 目的")
doc.add_paragraph("規範空壓機每日點檢項目與判定基準，及早發現異常，避免非計畫停機。")

h("2. 適用範圍")
doc.add_paragraph("適用於本廠所有螺旋式空壓機及附屬之冷凍式乾燥機。")

h("3. 權責")
doc.add_paragraph("3.1 設備保養人員：依本程序執行每日點檢並填寫點檢表。")
doc.add_paragraph("3.2 設備主管：每週覆核點檢表，異常項目應於 24 小時內安排維修。")

h("4. 安全注意事項")
doc.add_paragraph("4.1 點檢前應穿戴安全眼鏡、耳塞及防護手套。")
doc.add_paragraph("4.2 打開機殼或更換濾網前，必須執行上鎖掛牌，並確認系統壓力已洩放至 0 MPa。")
doc.add_paragraph("4.3 機體運轉中禁止觸碰排氣管路，表面溫度可能超過 80°C。")

h("5. 點檢項目與判定基準")
rows = [
    ("項次", "點檢項目", "判定基準", "異常處理"),
    ("1", "排氣壓力", "0.65～0.75 MPa", "超出範圍時通知設備主管"),
    ("2", "排氣溫度", "不超過 100°C", "清潔冷卻器並檢查潤滑油油位"),
    ("3", "潤滑油油位", "油位計 1/2～3/4 之間", "補充同型號潤滑油"),
    ("4", "空氣濾網壓差", "不超過 0.05 bar", "更換空氣濾網"),
    ("5", "冷凍式乾燥機露點", "3°C 以下", "檢查冷媒壓力與排水器"),
    ("6", "自動排水器", "每 30 分鐘正常排水", "清潔排水器濾網"),
    ("7", "運轉異音與振動", "無異音、無異常振動", "停機並通報維修"),
]
t = doc.add_table(rows=len(rows), cols=4)
t.style = "Table Grid"
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        t.cell(i, j).text = v

h("6. 作業步驟")
doc.add_paragraph("6.1 確認控制面板無警報顯示，記錄運轉時數。")
doc.add_paragraph("6.2 依第 5 節項目逐項點檢，並將量測值填入點檢表。")
doc.add_paragraph("6.3 發現異常時，於點檢表註記並立即通報設備主管。")
doc.add_paragraph("6.4 潤滑油每運轉 4,000 小時更換一次，油過濾器同時更換。")
doc.add_paragraph("6.5 點檢完成後，點檢表交由設備主管簽核並歸檔保存 1 年。")

h("7. 相關文件")
doc.add_paragraph("7.1 空壓機操作手冊")
doc.add_paragraph("7.2 上鎖掛牌管理程序（SOP-SF-003）")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("ok", OUT)
